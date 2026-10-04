"""
"Miya" — tayyor darslar bazasi (Excel → bot).

Excel tuzilmasi (Fanlar_Miya.xlsx):
  • Yoriqnoma — qo‘llanma (import qilinmaydi)
  • Mavzular  — ko‘rish uchun ro‘yxat (import qilinmaydi)
  • Miya      — darslar: har qator = bitta dars (sinf + 2 fan + mavzu)

Bot ishga tushganda data/Fanlar_Miya.xlsx fayli avtomatik yuklanadi (fayl o‘zgargan bo‘lsa).
Admin /miya_import orqali yangi Excel yuborib ham yangilashi mumkin.
"""
from __future__ import annotations

import hashlib
import io
import logging
import os
import re
from dataclasses import dataclass, field

from . import lesson as L
from . import objects, subjects

log = logging.getLogger(__name__)

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BUNDLED_PATH = os.path.join(BASE_DIR, "data", "Fanlar_Miya.xlsx")

COLUMNS = [
    "Sinf", "Fan_1", "Fan_2", "Mavzu", "Daqiqa", "Malumot", "Integratsiya", "Maqsad_1", "Maqsad_2",
    "Jihozlar", "Dars_rejasi", "Topshiriqlar", "Javoblar", "Uyga_vazifa", "Tekshirish", "Rasm_tavsifi",
    "Rasm_talablari", "Kalit_sozlar", "Korgazma_sarlavhasi", "Korgazma_ustunlari", "Korgazma_qatorlari",
    "Korgazma_izohi", "Manba",
]
REQUIRED = ["Sinf", "Fan_1", "Fan_2", "Mavzu", "Malumot", "Topshiriqlar", "Javoblar"]
MAX_ROWS = 1000
MAX_BYTES = 6 * 1024 * 1024
LEVELS = {"oson": "oson", "o'rta": "o‘rta", "orta": "o‘rta", "murakkab": "murakkab", "qiyin": "murakkab"}
STORY_HEADERS = ("O‘QISH UCHUN MUALLIFLIK HIKOYASI", "O‘QISH UCHUN MUALLIFLIK MATNI")

_SUBJECT_ALIASES = {
    "ona tili": "ona", "onatili": "ona",
    "oqish": "oqish", "oqish savodxonligi": "oqish", "adabiy oqish": "oqish",
    "matematika": "mat",
    "tabiiy fan": "tab", "tabiiy fanlar": "tab",
    "texnologiya": "tex",
}


class RowError(ValueError):
    pass


def _plain(text) -> str:
    s = str(text or "").strip().lower()
    s = re.sub(r"[‘’ʻʼ`´']", "", s)
    return re.sub(r"\s+", " ", s)


def subject_key(name) -> str | None:
    return _SUBJECT_ALIASES.get(_plain(name))


def pair_key_for(a, b) -> str | None:
    ka, kb_ = subject_key(a), subject_key(b)
    if not ka or not kb_:
        return None
    for key, x, y in subjects.PAIRS:
        if {x, y} == {ka, kb_}:
            return key
    return None


def topic_key(topic: str) -> str:
    s = _plain(topic)
    s = re.sub(r"[«»\"“”„.,!?;:()\-—–]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def _lines(value) -> list[str]:
    return [line.strip(" •-\t") for line in str(value or "").replace("\r", "").split("\n") if line.strip(" •-\t")]


def _pipe(line: str, n: int) -> list[str]:
    parts = [p.strip() for p in line.split("|")]
    if len(parts) > n:
        parts = parts[: n - 1] + [" | ".join(parts[n - 1:])]
    return parts + [""] * (n - len(parts))


def _int(value, default=None):
    if value is None or str(value).strip() == "":
        return default
    try:
        return int(float(str(value).strip()))
    except ValueError:
        return default


def story_part(info: str) -> str:
    """Ma’lumot matnidan bolalarga o‘qib beriladigan qismni (hikoya yoki vaziyat) ajratadi."""
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", info or "") if p.strip()]
    if not paragraphs:
        return ""
    first = paragraphs[0]
    lines = first.split("\n")
    if lines and lines[0].strip().upper().replace("'", "‘") in STORY_HEADERS:
        first = "\n".join(lines[1:]).strip()
    return first


def exhibit_groups(rows: list[list[str]]) -> list[dict]:
    """Ko‘rgazma qatorlari "6 ta olma" ko‘rinishida bo‘lsa — hisob kartasi guruhlari."""
    groups = []
    for label, value in rows:
        m = re.match(r"^\s*(\d[\d\s]*)\s*ta\s+(.+?)\s*\.?$", value or "")
        if not m:
            return []
        count = int(re.sub(r"\s", "", m.group(1)))
        name = m.group(2).strip()
        key = objects.guess_key(name)
        if not key or count < 1:
            return []
        groups.append({"label": label, "where_en": "", "object_key": key, "object_uz": name, "count": count})
    return groups


def row_to_lesson(row: dict) -> dict:
    """Excel qatorini bot ichidagi dars tuzilmasiga aylantiradi. Xato bo‘lsa RowError."""
    grade = _int(row.get("Sinf"))
    if grade not in subjects.GRADES:
        raise RowError(f"Sinf 1–4 bo‘lishi kerak ({row.get('Sinf')!r})")
    pair = pair_key_for(row.get("Fan_1"), row.get("Fan_2"))
    if not pair:
        raise RowError(f"fan juftligi noma’lum: {row.get('Fan_1')!r} + {row.get('Fan_2')!r}")
    topic = L.uz_fix(row.get("Mavzu"))
    if not 2 <= len(topic) <= 150:
        raise RowError("Mavzu bo‘sh yoki juda uzun")
    minutes = _int(row.get("Daqiqa"), 45)
    if minutes not in (30, 35, 40, 45, 60, 80, 90):
        raise RowError(f"Daqiqa noto‘g‘ri ({row.get('Daqiqa')!r})")
    info = L.uz_fix(row.get("Malumot"))
    if len(info) < 20:
        raise RowError("Malumot juda qisqa")

    answers: dict[int, tuple[str, str]] = {}
    for line in _lines(row.get("Javoblar")):
        num, text, crit = _pipe(line, 3)
        n = _int(num)
        if n is not None:
            answers[n] = (L.uz_fix(text), L.uz_fix(crit))
    tasks = []
    for line in _lines(row.get("Topshiriqlar")):
        num, level, uses, question = _pipe(line, 4)
        n = _int(num)
        if n is None or not question:
            raise RowError(f"Topshiriq qatori noto‘g‘ri: {line[:60]!r}")
        ans, crit = answers.get(n, ("", ""))
        tasks.append({
            "title": f"{n}-topshiriq",
            "subjects": subjects.pair_label(pair),
            "instruction": L.uz_fix(question),
            "uses_image": str(uses).strip() in ("1", "ha", "yes", "true"),
            "answer": ans,
            "criteria": crit,
            "level": LEVELS.get(_plain(level).replace("‘", "'"), L.uz_fix(level)),
            "check_expression": "",
            "check_result": "",
        })
    if not 3 <= len(tasks) <= 12:
        raise RowError(f"topshiriqlar soni 3–12 bo‘lishi kerak (bor: {len(tasks)})")
    missing = [i for i in range(1, len(tasks) + 1) if i not in answers]
    if missing:
        raise RowError(f"javobi yo‘q topshiriqlar: {missing}")

    plan = []
    for line in _lines(row.get("Dars_rejasi")):
        mins, stage, teacher, students = _pipe(line, 4)
        m = _int(mins)
        if m is None or not stage:
            continue
        plan.append({"stage": L.uz_fix(stage), "minutes": m, "teacher": L.uz_fix(teacher), "students": L.uz_fix(students)})
    L.fix_minutes(plan, minutes)

    columns = _pipe(str(row.get("Korgazma_ustunlari") or "Belgi | Ma’lumot"), 2)
    ex_rows = []
    for line in _lines(row.get("Korgazma_qatorlari")):
        a, b = _pipe(line, 2)
        if a or b:
            ex_rows.append([L.uz_fix(a), L.uz_fix(b)])
    exhibit = None
    if ex_rows:
        exhibit = {
            "title": L.uz_fix(row.get("Korgazma_sarlavhasi")) or topic,
            "columns": [L.uz_fix(columns[0]) or "Belgi", L.uz_fix(columns[1]) or "Ma’lumot"],
            "rows": ex_rows[:12],
            "note": L.uz_fix(row.get("Korgazma_izohi")),
        }

    criteria = []
    for _, crit in answers.values():
        if crit and crit not in criteria:
            criteria.append(crit)
    story = story_part(info)
    image_req = [L.uz_fix(x) for x in _lines(row.get("Rasm_talablari"))]
    scene = str(row.get("Rasm_tavsifi") or "").strip()

    lesson = {
        "title": topic,
        "summary": "",
        "objectives": {
            "subject_a": [L.uz_fix(row.get("Maqsad_1"))] if row.get("Maqsad_1") else [],
            "subject_b": [L.uz_fix(row.get("Maqsad_2"))] if row.get("Maqsad_2") else [],
            "tarbiyaviy": [],
        },
        "teacher_info": info,
        "key_points": [],
        "child_text": story,
        "integration": {"idea": L.uz_fix(row.get("Integratsiya")), "subject_a_part": "", "subject_b_part": ""},
        "vocabulary": [],
        "materials": [L.uz_fix(x) for x in _lines(row.get("Jihozlar"))],
        "image": {
            "scene_en": scene,
            "description_uz": "",
            "count_groups": exhibit_groups(exhibit["rows"]) if exhibit else [],
            "questions": [],
        },
        "tasks": tasks,
        "lesson_plan": plan,
        "assessment": criteria,
        "differentiation": {"support": "", "challenge": ""},
        "homework": L.uz_fix(row.get("Uyga_vazifa")),
        "voice_text": story,
        "meta": {"grade": grade, "pair": pair, "topic": topic, "minutes": minutes},
        "origin": "miya",
        "exhibit": exhibit,
        "keywords": [L.uz_fix(x) for x in _lines(row.get("Kalit_sozlar")) if _plain(x) != _plain(topic)][:10],
        "teacher_notes": [L.uz_fix(x) for x in _lines(row.get("Tekshirish"))][:8],
        "image_requirements": image_req[:6],
        "source": L.uz_fix(row.get("Manba")),
    }
    L.check_math(lesson)
    return lesson


@dataclass
class ParseResult:
    lessons: list[dict] = field(default_factory=list)
    raws: list[dict] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    total_rows: int = 0


def _header_map(header: list) -> dict[str, int]:
    wanted = {_plain(c).replace(" ", "_"): c for c in COLUMNS}
    mapping = {}
    for idx, value in enumerate(header):
        key = _plain(value).replace(" ", "_")
        if key in wanted:
            mapping[wanted[key]] = idx
    return mapping


def parse_workbook(data: bytes) -> ParseResult:
    import openpyxl

    result = ParseResult()
    if len(data) > MAX_BYTES:
        result.errors.append("Fayl hajmi 6 MB dan katta.")
        return result
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        result.errors.append(f"Excel fayl ochilmadi: {exc}")
        return result
    try:
        sheet = next((ws for ws in wb.worksheets if _plain(ws.title) == "miya"), None)
        if sheet is None:
            result.errors.append("«Miya» nomli varaq topilmadi.")
            return result
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None) or []
        mapping = _header_map(list(header))
        missing = [c for c in REQUIRED if c not in mapping]
        if missing:
            result.errors.append("Sarlavhada ustunlar yetishmaydi: " + ", ".join(missing))
            return result
        seen: dict[tuple, int] = {}
        for excel_row, values in enumerate(rows, start=2):
            values = list(values or [])
            if not any(v not in (None, "") for v in values):
                continue
            result.total_rows += 1
            if result.total_rows > MAX_ROWS:
                result.errors.append(f"Bir importda {MAX_ROWS} tadan ko‘p dars bo‘lmasin.")
                break
            raw = {col: (values[idx] if idx < len(values) else None) for col, idx in mapping.items()}
            try:
                lesson = row_to_lesson(raw)
            except RowError as exc:
                result.errors.append(f"{excel_row}-qator: {exc}")
                continue
            meta = lesson["meta"]
            key = (meta["grade"], meta["pair"], topic_key(meta["topic"]))
            if key in seen:
                result.warnings.append(f"{excel_row}-qator: {seen[key]}-qator bilan bir xil dars — keyingisi olinadi.")
                idx = [i for i, l in enumerate(result.lessons)
                       if (l["meta"]["grade"], l["meta"]["pair"], topic_key(l["meta"]["topic"])) == key][0]
                result.lessons[idx] = lesson
                result.raws[idx] = {k: ("" if v is None else v) for k, v in raw.items()}
                continue
            seen[key] = excel_row
            problems = L.math_problems(lesson)
            if problems:
                result.warnings.append(f"{excel_row}-qator ({meta['topic']}, {meta['grade']}-sinf): "
                                       f"{', '.join(map(str, problems))}-javobda hisob xatosi bo‘lishi mumkin.")
            result.lessons.append(lesson)
            result.raws.append({k: ("" if v is None else v) for k, v in raw.items()})
    finally:
        wb.close()
    return result


def file_hash(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()[:16]


# ───────────────────────── Bazaga yozish ─────────────────────────

def import_lessons(db, result: ParseResult, source: str) -> tuple[int, int]:
    """Darslarni bazaga yozadi (bir xil sinf+fanlar+mavzu yangilanadi). Qaytaradi: (yangi, yangilangan)."""
    new = updated = 0
    for lesson, raw in zip(result.lessons, result.raws):
        meta = lesson["meta"]
        state = db.upsert_catalog(meta["grade"], meta["pair"], meta["topic"], topic_key(meta["topic"]),
                                  meta["minutes"], lesson, raw, source)
        if state == "new":
            new += 1
        elif state == "updated":
            updated += 1
    return new, updated


def seed_from_bundle(db, path: str = BUNDLED_PATH, force: bool = False) -> str:
    """Git bilan kelgan Excel faylni (o‘zgargan bo‘lsa) bazaga yuklaydi. Natija matnini qaytaradi."""
    if not os.path.exists(path):
        return "Miya fayli topilmadi: " + path
    with open(path, "rb") as fh:
        data = fh.read()
    digest = file_hash(data)
    if not force and db.get_setting("miya_seed_hash") == digest and db.count_catalog() > 0:
        return f"Miya o‘zgarmagan ({db.count_catalog()} ta dars)."
    result = parse_workbook(data)
    if result.errors and not result.lessons:
        log.error("Miya fayli o‘qilmadi: %s", result.errors[:3])
        return "Miya fayli o‘qilmadi: " + "; ".join(result.errors[:3])
    new, updated = import_lessons(db, result, "git")
    db.set_setting("miya_seed_hash", digest)
    msg = f"Miya yuklandi: {len(result.lessons)} ta dars (yangi {new}, yangilangan {updated})."
    if result.errors:
        msg += f" Xato qatorlar: {len(result.errors)}."
    return msg


# ───────────────────────── Excelga eksport ─────────────────────────

def _join_lines(items) -> str:
    return "\n".join(str(x) for x in items if str(x).strip())


def lesson_to_row(lesson: dict) -> dict:
    """Bot ichidagi darsni (masalan AI darsini) Miya qatoriga aylantiradi."""
    meta = lesson["meta"]
    a, b = subjects.pair_subjects(meta["pair"])
    obj = lesson.get("objectives", {})
    tasks = lesson.get("tasks", [])
    exhibit = lesson.get("exhibit")
    if exhibit:
        ex_title, ex_cols = exhibit["title"], " | ".join(exhibit["columns"])
        ex_rows = _join_lines(f"{r[0]} | {r[1]}" for r in exhibit["rows"])
        ex_note = exhibit.get("note", "")
    else:
        groups = lesson.get("image", {}).get("count_groups", [])
        ex_title, ex_cols = lesson["title"], "Belgi | Ma’lumot"
        ex_rows = _join_lines(f"{g['label']} | {g['count']} ta {g['object_uz']}" for g in groups)
        ex_note = ""
    info = lesson.get("teacher_info", "")
    if lesson.get("origin") != "miya" and lesson.get("child_text"):
        info = f"{lesson['child_text']}\n\n{info}".strip()
    return {
        "Sinf": meta["grade"],
        "Fan_1": subjects.subject_name(a),
        "Fan_2": subjects.subject_name(b),
        "Mavzu": meta["topic"] or lesson["title"],
        "Daqiqa": meta["minutes"],
        "Malumot": info,
        "Integratsiya": " ".join(x for x in (lesson.get("integration") or {}).values() if x),
        "Maqsad_1": "; ".join(obj.get("subject_a", [])),
        "Maqsad_2": "; ".join(obj.get("subject_b", [])),
        "Jihozlar": _join_lines(lesson.get("materials", [])),
        "Dars_rejasi": _join_lines(f"{s['minutes']} | {s['stage']} | {s['teacher']} | {s['students']}"
                                   for s in lesson.get("lesson_plan", [])),
        "Topshiriqlar": _join_lines(f"{i} | {t.get('level') or 'o‘rta'} | {1 if t.get('uses_image') else 0} | "
                                    f"{t['instruction']}" for i, t in enumerate(tasks, 1)),
        "Javoblar": _join_lines(f"{i} | {t.get('answer') or '—'} | {t.get('criteria') or ''}".rstrip(" |")
                                for i, t in enumerate(tasks, 1)),
        "Uyga_vazifa": lesson.get("homework", ""),
        "Tekshirish": _join_lines(lesson.get("teacher_notes", [])),
        "Rasm_tavsifi": lesson.get("image", {}).get("scene_en", ""),
        "Rasm_talablari": _join_lines(lesson.get("image_requirements", [])),
        "Kalit_sozlar": _join_lines([meta["topic"]] + list(lesson.get("keywords", []))),
        "Korgazma_sarlavhasi": ex_title,
        "Korgazma_ustunlari": ex_cols,
        "Korgazma_qatorlari": ex_rows,
        "Korgazma_izohi": ex_note,
        "Manba": lesson.get("source", "") or ("AI yordamida tayyorlangan dars." if lesson.get("origin") == "ai" else ""),
    }


YORIQNOMA = [
    ("Nima tayyor?", "Har bir qator — bitta tayyor dars: sinf (1–4) + ikki fan + mavzu. Bot bu darslarni AI so‘rovisiz, "
                     "darhol ochadi. Bular tavsiya etilgan mualliflik namunalari, rasmiy taqvim-mavzu reja emas."),
    ("Botdagi tartib", "📚 Dars tayyorlash → sinf → fan juftligi → 📗 tayyor mavzu. Ma’lumot, ko‘rgazma kartasi, "
                       "topshiriqlar, javoblar, dars rejasi, Word fayl va ovozli o‘qish darhol ochiladi."),
    ("Git orqali yangilash", "Shu faylni loyihadagi data/Fanlar_Miya.xlsx o‘rniga qo‘ying va GitHub’ga yuklang. "
                             "Bot qayta ishga tushganda fayl o‘zgarganini sezadi va miyani avtomatik to‘ldiradi."),
    ("Bot orqali yangilash", "Admin /miya_import yozadi → .xlsx faylni yuboradi → tekshiruv natijasini ko‘radi → "
                             "tasdiqlaydi. Bir xil sinf + fanlar + mavzu yangilanadi; qolgan darslar o‘chmaydi."),
    ("Qaysi varaq import qilinadi?", "Faqat «Miya» varag‘i. «Mavzular» — ko‘rish uchun ro‘yxat. «Miya» varag‘i nomi "
                                     "va birinchi qator sarlavhalarini o‘zgartirmang."),
    ("Ro‘yxatli kataklar", "Jihozlar, Tekshirish, Rasm_talablari, Kalit_sozlar: har band alohida qatorda "
                           "(Excelda yangi qator: Alt+Enter)."),
    ("Dars_rejasi", "Har bosqich alohida qatorda: daqiqa | nomi | o‘qituvchi harakati | o‘quvchi harakati. "
                    "Daqiqalar yig‘indisi Daqiqa ustuniga teng bo‘lsin (bo‘lmasa bot o‘zi moslaydi)."),
    ("Topshiriqlar", "Har mashq: raqam | daraja | ko‘rgazma | savol. Daraja: oson, o‘rta yoki murakkab. "
                     "Ko‘rgazma: 1 = kartadan foydalanadi, 0 = foydalanmaydi. 3–12 ta mashq."),
    ("Javoblar", "Har javob: raqam | javob matni | baholash mezoni. Raqamlar topshiriq raqamlariga mos bo‘lsin. "
                 "Hisoblar «6+4=10» ko‘rinishida yozilsa, bot ularni o‘zi tekshiradi."),
    ("Ko‘rgazma", "Korgazma_ustunlari: birinchi | ikkinchi sarlavha. Korgazma_qatorlari: har qator «belgi | ma’lumot». "
                  "Ma’lumot «6 ta olma» ko‘rinishida bo‘lsa, bot rasmli hisob kartasini chizadi; aks holda jadval-karta."),
    ("Rasm", "Ko‘rgazma kartasi doim aniq (bot o‘zi chizadi). AI rasm ixtiyoriy: Rasm_tavsifi asosida chiziladi, "
             "API sarfi bor va bir marta chizilgach hamma uchun saqlanib qoladi."),
    ("Cheklovlar", "Sinf: 1–4. Daqiqa: 30, 35, 40, 45, 60, 80 yoki 90. Fayl: 6 MB gacha; bir importda 1000 darsgacha. "
                   "Formula va makros ishlatmang."),
]


def build_workbook(rows: list[dict]) -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    teal = PatternFill("solid", fgColor="154F60")
    light = PatternFill("solid", fgColor="E4F2F1")
    white_bold = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    thin = Side(style="thin", color="C9D6D8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap_top = Alignment(wrap_text=True, vertical="top")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Yoriqnoma"
    for col in "ABCDEF":
        ws.column_dimensions[col].width = 18
    ws.merge_cells("A1:F2")
    ws["A1"] = "FANLAR INTEGRATSIYASI | TAYYOR MIYA"
    ws["A1"].font = Font(name="Calibri", bold=True, size=18, color="FFFFFF")
    ws["A1"].fill = teal
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    tasks_total = sum(len(_lines(r.get("Topshiriqlar"))) for r in rows)
    pairs = {(r.get("Fan_1"), r.get("Fan_2")) for r in rows}
    for col, (title, value) in zip(("A", "C", "E"), (("Tayyor darslar", len(rows)), ("Topshiriqlar", tasks_total),
                                                      ("Fan juftliklari", len(pairs)))):
        end = chr(ord(col) + 1)
        ws.merge_cells(f"{col}4:{end}4")
        ws.merge_cells(f"{col}5:{end}6")
        ws[f"{col}4"] = title
        ws[f"{col}5"] = value
        for ref in (f"{col}4", f"{col}5"):
            ws[ref].fill = light
            ws[ref].alignment = Alignment(horizontal="center", vertical="center")
        ws[f"{col}5"].font = Font(name="Calibri", bold=True, size=22, color="154F60")
    ws.merge_cells("A8:B8")
    ws.merge_cells("C8:F8")
    ws["A8"], ws["C8"] = "BO‘LIM", "QO‘LLASH TARTIBI"
    for ref in ("A8", "C8"):
        ws[ref].font = white_bold
        ws[ref].fill = teal
    for i, (title, text) in enumerate(YORIQNOMA, start=9):
        ws.merge_cells(f"A{i}:B{i}")
        ws.merge_cells(f"C{i}:F{i}")
        ws[f"A{i}"] = title
        ws[f"C{i}"] = text
        ws[f"A{i}"].font = Font(name="Calibri", bold=True)
        ws[f"A{i}"].alignment = wrap_top
        ws[f"C{i}"].alignment = wrap_top
        ws.row_dimensions[i].height = 64

    ws2 = wb.create_sheet("Mavzular")
    head = ["№", "Sinf", "Fan 1", "Fan 2", "Mavzu", "Daqiqa", "Topshiriqlar", "Material holati"]
    ws2.append(head)
    for i, r in enumerate(rows, start=1):
        ws2.append([i, r.get("Sinf"), r.get("Fan_1"), r.get("Fan_2"), r.get("Mavzu"), r.get("Daqiqa"),
                    len(_lines(r.get("Topshiriqlar"))), "Mualliflik namuna"])
    for col, width in zip("ABCDEFGH", (6, 8, 18, 18, 38, 10, 15, 23)):
        ws2.column_dimensions[col].width = width
    for cell in ws2[1]:
        cell.font = white_bold
        cell.fill = teal
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws2.freeze_panes = "A2"

    ws3 = wb.create_sheet("Miya")
    ws3.append(COLUMNS)
    for r in rows:
        ws3.append([r.get(c, "") for c in COLUMNS])
    widths = [7, 16, 16, 34, 9] + [38] * 12 + [24, 30, 24, 38, 38, 38]
    for idx, width in enumerate(widths, start=1):
        ws3.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = width
    for cell in ws3[1]:
        cell.font = white_bold
        cell.fill = teal
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws3.row_dimensions[1].height = 36
    for row in ws3.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap_top
            cell.border = border
    ws3.freeze_panes = "E2"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
