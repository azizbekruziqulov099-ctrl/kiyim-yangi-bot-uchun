"""Darsni Word (.docx) fayl — "dars ishlanmasi" ko‘rinishida tayyorlash."""
from __future__ import annotations

import io
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import lesson as L
from . import subjects

FONT = "Times New Roman"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)


def _setup(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(1.5)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.1
    for name in ("Heading 1", "Heading 2", "Title"):
        style = doc.styles[name]
        style.font.name = FONT
        rfonts = style.element.rPr.rFonts
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rfonts.attrib.pop(qn(attr), None)
        rfonts.set(qn("w:eastAsia"), FONT)
        style.font.color.rgb = ACCENT
    doc.styles["Heading 1"].font.size = Pt(14)
    doc.styles["Heading 2"].font.size = Pt(12.5)
    for name in ("List Bullet", "List Number"):
        doc.styles[name].font.name = FONT
        doc.styles[name].font.size = Pt(12)


def _heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)


def _para(doc: Document, text: str, bold: bool = False, italic: bool = False, size: float | None = None,
          align=None) -> None:
    if not text:
        return
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    if size:
        run.font.size = Pt(size)
    if align is not None:
        p.alignment = align


def _bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        if item:
            doc.add_paragraph(item, style="List Bullet")


def _label_para(doc: Document, label: str, text: str) -> None:
    if not text:
        return
    p = doc.add_paragraph()
    p.add_run(label).bold = True
    p.add_run(" " + text)


def _set_widths(table, widths) -> None:
    table.autofit = False
    for i, width in enumerate(widths):
        table.columns[i].width = width
        for cell in table.columns[i].cells:
            cell.width = width


def _cell(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text or "")
    run.bold = bold
    run.font.size = Pt(11)


def build_docx(lesson: dict, image: bytes | None = None, card: bytes | None = None,
               author: str = "", image_note: str = "") -> bytes:
    meta = lesson["meta"]
    a, b = subjects.pair_subjects(meta["pair"])
    doc = Document()
    _setup(doc)

    _para(doc, "DARS ISHLANMASI", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(lesson["title"])
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = ACCENT
    _para(doc, "Integratsiyalashgan dars", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    rows = [
        ("Sinf", f"{meta['grade']}-sinf"),
        ("Fanlar", subjects.pair_label(meta["pair"])),
        ("Mavzu", meta["topic"] or lesson["title"]),
        ("Dars turi", "Fanlararo integratsiyalashgan dars"),
        ("Davomiyligi", f"{meta['minutes']} daqiqa"),
        ("O‘qituvchi", author or "______________________"),
        ("Sana", "______________________"),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate(rows):
        _cell(table.cell(i, 0), k, bold=True)
        _cell(table.cell(i, 1), v)
    _set_widths(table, (Cm(4), Cm(12.5)))

    if lesson["summary"]:
        doc.add_paragraph()
        _para(doc, lesson["summary"], italic=True)

    obj = lesson["objectives"]
    _heading(doc, "Dars maqsadlari")
    if obj["subject_a"]:
        _para(doc, f"{subjects.subject_name(a)}:", bold=True)
        _bullets(doc, obj["subject_a"])
    if obj["subject_b"]:
        _para(doc, f"{subjects.subject_name(b)}:", bold=True)
        _bullets(doc, obj["subject_b"])
    if obj["tarbiyaviy"]:
        _para(doc, "Tarbiyaviy:", bold=True)
        _bullets(doc, obj["tarbiyaviy"])

    ready = lesson.get("origin") == "miya"
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", lesson.get("teacher_info") or "") if p.strip()]
    first = paragraphs[0].split("\n") if paragraphs else [""]
    if first[0].strip().upper() in L.STORY_HEADERS:
        _heading(doc, "O‘qish uchun matn")
        _para(doc, "\n".join(first[1:]).strip(), italic=True)
        if paragraphs[1:]:
            _heading(doc, "O‘qituvchiga", level=2)
            for para in paragraphs[1:]:
                _para(doc, para)
    elif paragraphs or lesson.get("key_points"):
        _heading(doc, "Mavzu bo‘yicha ma’lumot")
        for para in paragraphs:
            _para(doc, para)
        _bullets(doc, lesson.get("key_points", []))
    if lesson.get("keywords"):
        _label_para(doc, "Kalit so‘zlar:", ", ".join(lesson["keywords"]))

    if lesson.get("child_text") and not ready:
        _heading(doc, "O‘quvchilar uchun matn")
        _para(doc, lesson["child_text"], italic=True)

    integ = lesson["integration"]
    if any(integ.values()):
        _heading(doc, "Fanlar qanday bog‘lanadi")
        _para(doc, integ["idea"])
        _label_para(doc, f"{subjects.subject_name(a)}:", integ["subject_a_part"])
        _label_para(doc, f"{subjects.subject_name(b)}:", integ["subject_b_part"])

    if lesson["vocabulary"]:
        _heading(doc, "Yangi so‘zlar")
        _bullets(doc, [f"{v['word']} — {v['meaning']}" if v["meaning"] else v["word"] for v in lesson["vocabulary"]])

    if lesson["materials"]:
        _heading(doc, "Jihozlar")
        _bullets(doc, lesson["materials"])

    if image:
        _heading(doc, "Dars rasmi")
        doc.add_picture(io.BytesIO(image), width=Cm(16))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _para(doc, lesson["image"]["description_uz"], italic=True)
        if image_note:
            _para(doc, re.sub(r"<[^>]+>", "", image_note), size=10.5, italic=True)
    if lesson["image"]["questions"]:
        _para(doc, "Rasm bilan ishlash savollari:", bold=True)
        for q in lesson["image"]["questions"]:
            doc.add_paragraph(q, style="List Number")

    if card:
        _heading(doc, "Ko‘rgazma kartasi" if lesson.get("exhibit") else "Hisob kartasi (aniq sonlar)")
        doc.add_picture(io.BytesIO(card), width=Cm(15))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _bullets(doc, L.count_lines(lesson))

    if lesson["lesson_plan"]:
        _heading(doc, "Dars borishi")
        plan = doc.add_table(rows=1, cols=5)
        plan.style = "Table Grid"
        widths = (Cm(0.9), Cm(3.2), Cm(1.6), Cm(5.6), Cm(5.2))
        for cell, text in zip(plan.rows[0].cells, ("№", "Bosqich", "Vaqt", "O‘qituvchi faoliyati",
                                                   "O‘quvchi faoliyati")):
            _cell(cell, text, bold=True)
        for i, st in enumerate(lesson["lesson_plan"], start=1):
            cells = plan.add_row().cells
            for cell, text in zip(cells, (str(i), st["stage"], f"{st['minutes']} daq.", st["teacher"],
                                          st["students"])):
                _cell(cell, text)
        _set_widths(plan, widths)
        total = sum(st["minutes"] for st in lesson["lesson_plan"])
        _para(doc, f"Jami: {total} daqiqa", italic=True, size=10.5)

    _heading(doc, "Integratsiyalashgan topshiriqlar")
    for i, t in enumerate(lesson["tasks"], start=1):
        p = doc.add_paragraph()
        if ready:
            p.add_run(f"{i}-topshiriq").bold = True
            note = f"  — {t.get('level') or ''}" + (", kartadan" if t.get("uses_image") else "")
        else:
            p.add_run(f"{i}. {t['title']}").bold = True
            note = f"  — {t['subjects']}" + (" (rasm bilan)" if t.get("uses_image") else "")
        r = p.add_run(note)
        r.italic = True
        r.font.size = Pt(10.5)
        _para(doc, t["instruction"])

    if lesson["assessment"]:
        _heading(doc, "Baholash mezonlari")
        _bullets(doc, lesson["assessment"])
    diff = lesson["differentiation"]
    if diff["support"] or diff["challenge"]:
        _heading(doc, "Tabaqalashtirish")
        _label_para(doc, "Qiynalayotgan o‘quvchilarga:", diff["support"])
        _label_para(doc, "Kuchli o‘quvchilarga:", diff["challenge"])
    if lesson["homework"]:
        _heading(doc, "Uyga vazifa")
        _para(doc, lesson["homework"])
    if lesson.get("teacher_notes"):
        _heading(doc, "O‘qituvchi uchun eslatma")
        _bullets(doc, lesson["teacher_notes"])

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    _heading(doc, "Javoblar (o‘qituvchi uchun)")
    for i, t in enumerate(lesson["tasks"], start=1):
        p = doc.add_paragraph()
        p.add_run(f"{i}-topshiriq: " if ready else f"{i}. {t['title']}: ").bold = True
        p.add_run(t["answer"] or "—")
        for chk in t.get("checks") or []:
            if chk["ok"] is False:
                text = (f"Tekshiruv: {chk['expr']} = {chk['value']}, javobda «{chk['stated']}» yozilgan — "
                        f"tekshirib oling.")
            else:
                text = f"Tekshiruv: {chk['expr']} = {chk['value']} ✓"
            _para(doc, text, italic=True, size=10.5)

    doc.add_paragraph()
    if lesson.get("source"):
        _para(doc, lesson["source"], italic=True, size=9.5)
    _para(doc, "Material «Fanlar integratsiyasi» boti yordamida tayyorlandi. Darsdan oldin o‘qituvchi "
               "mazmun, sonlar va rasmni tekshirib chiqishi tavsiya etiladi.", italic=True, size=9.5)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def file_name(lesson: dict) -> str:
    topic = lesson["meta"].get("topic") or lesson["title"]
    slug = re.sub(r"[‘’ʻʼ'`]", "", topic)
    slug = re.sub(r"[^0-9A-Za-z]+", "_", slug).strip("_")[:40] or "dars"
    return f"Dars_{lesson['meta']['grade']}-sinf_{slug}.docx"
