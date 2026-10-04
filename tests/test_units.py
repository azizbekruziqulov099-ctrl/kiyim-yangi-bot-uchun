"""Alohida qismlar sinovi: matematik tekshiruv, miya, baza, AI ulanishi, kartalar, Word, ovoz, sozlamalar."""
import asyncio
import io
import json
from fractions import Fraction

import httpx
import pytest
from docx import Document
from PIL import Image

from edu import card, docx_export, miya, objects, prompts, tts
from edu import lesson as L
from edu.ai import AIError, AIService, parse_json_loose
from edu.config import load_config

from .conftest import make_cfg
from .fake_ai import openai_transport


# ───────── matn va matematika ─────────

def test_uz_fix():
    assert L.uz_fix("O'zbekiston bog'lari, ta'lim, she'r, mo`jiza") == "O‘zbekiston bog‘lari, ta’lim, she’r, mo‘jiza"
    assert L.uz_fix("  ikki   bo'shliq \n\n\n\n uch") == "ikki bo‘shliq\n\nuch"


@pytest.mark.parametrize("expr,val", [("6+4", 10), ("(360+240):4", 150), ("3x4", 12), ("2,5*2", 5),
                                      ("1 260 + 940", 2200), ("15÷3", 5), ("10 − 4", 6), ("6+4=10", 10)])
def test_safe_eval(expr, val):
    assert L.safe_eval(expr) == Fraction(val)


@pytest.mark.parametrize("expr", ["7/0", "__import__('os')", "2**100", "abc", ""])
def test_safe_eval_rejects(expr):
    assert L.safe_eval(expr) is None


def test_find_equations():
    eqs = L.find_equations("6+4=10. Jami 10 ta. (360+240):4=150; 600-150=450 ta qoldi, 18×5=90 cm²")
    assert eqs == [("6+4", "10"), ("(360+240):4", "150"), ("600-150", "450"), ("18×5", "90")]
    assert L.find_equations("Javob: 7 ta olma") == []


def test_check_math_flags_wrong_answer(sample_lesson):
    lesson = L.normalize_lesson(sample_lesson, grade=2, pair_key="ona_mat", topic="Bog'dagi hosil", minutes=45)
    assert L.math_problems(lesson) == [5]
    assert "⚠️" in L.fmt_answers(lesson)
    assert sum(s["minutes"] for s in lesson["lesson_plan"]) == 45


def test_normalize_handles_garbage():
    with pytest.raises(L.LessonError):
        L.normalize_lesson({"title": "x"}, grade=1, pair_key="ona_mat", topic="x", minutes=45)
    with pytest.raises(L.LessonError):
        L.normalize_lesson(["not", "dict"], grade=1, pair_key="ona_mat", topic="x", minutes=45)


def test_fix_minutes():
    plan = [{"minutes": 10}, {"minutes": 10}, {"minutes": 10}]
    assert sum(p["minutes"] for p in L.fix_minutes(plan, 45)) == 45
    plan = [{"minutes": 0}, {"minutes": 0}]
    assert sum(p["minutes"] for p in L.fix_minutes(plan, 90)) == 90


def test_split_message_respects_limit():
    text = "\n\n".join(f"<b>Bo‘lim {i}</b>\n" + "matn " * 300 for i in range(10))
    chunks = L.split_message(text)
    assert all(len(c) <= 4000 for c in chunks) and len(chunks) > 1
    assert all(c.count("<b>") == c.count("</b>") for c in chunks)


def test_cyr_to_lat():
    assert L.cyr_to_lat("Ўзбекистон мактаби") == "O‘zbekiston maktabi"
    assert L.cyr_to_lat("Ер юзи") == "Yer yuzi"
    assert L.cyr_to_lat("Lotin matn") == "Lotin matn"


def test_parse_json_loose():
    assert parse_json_loose('```json\n{"a": 1}\n```') == {"a": 1}
    assert parse_json_loose('Mana: {"a": 2} tamom') == {"a": 2}
    with pytest.raises(AIError):
        parse_json_loose("JSON emas")


# ───────── miya ─────────

def test_bundled_miya_is_complete(bundled_result):
    assert len(bundled_result.lessons) == 216
    assert not bundled_result.errors and not bundled_result.warnings
    from collections import Counter
    per = Counter((l["meta"]["pair"], l["meta"]["grade"]) for l in bundled_result.lessons)
    assert len(per) == 36 and set(per.values()) == {6}
    for lesson in bundled_result.lessons:
        assert sum(s["minutes"] for s in lesson["lesson_plan"]) == lesson["meta"]["minutes"]
        assert len(lesson["tasks"]) >= 6 and all(t["answer"] for t in lesson["tasks"])
        assert L.math_problems(lesson) == [], lesson["title"]
        text = json.dumps(lesson, ensure_ascii=False)
        if lesson["meta"]["grade"] > 1:
            assert "1-sinfga og‘zaki izoh" not in text


def test_miya_row_errors():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Miya"
    ws.append(miya.COLUMNS)
    ws.append([7, "Ona tili", "Matematika", "Xato sinf", 45, "Uzun ma’lumot matni bu yerda", "", "", "", "", "",
               "1 | oson | 0 | Savol", "1 | Javob | mezon"] + [""] * 10)
    ws.append([2, "Rus tili", "Matematika", "Noma’lum fan", 45, "Uzun ma’lumot matni bu yerda"] + [""] * 17)
    buf = io.BytesIO()
    wb.save(buf)
    result = miya.parse_workbook(buf.getvalue())
    assert not result.lessons and len(result.errors) == 2
    assert miya.parse_workbook(b"not excel").errors


def test_topic_key_and_pairs():
    assert miya.topic_key("Bog'dagi hosil!") == miya.topic_key("bog‘dagi  hosil") == "bogdagi hosil"
    assert miya.pair_key_for("Matematika", "Ona tili") == "ona_mat"
    assert miya.pair_key_for("O'qish", "Texnologiya") == "oqish_tex"
    assert miya.pair_key_for("Rus tili", "Matematika") is None


def test_exhibit_groups():
    assert miya.exhibit_groups([["1-savat", "6 ta olma"], ["2-savat", "4 ta olma"]])[0]["object_key"] == "apple"
    assert miya.exhibit_groups([["Qahramon", "Dilorom"]]) == []


def test_seed_from_bundle_only_when_changed(db):
    first = miya.seed_from_bundle(db)
    assert "216" in first
    assert "o‘zgarmagan" in miya.seed_from_bundle(db)
    assert db.count_catalog() == 216


def test_export_roundtrip(bundled_result):
    data = miya.build_workbook(bundled_result.raws)
    again = miya.parse_workbook(data)
    assert len(again.lessons) == 216 and not again.errors


# ───────── baza ─────────

def test_quota_and_refund(db):
    assert db.try_consume(7, "lessons", 2) and db.try_consume(7, "lessons", 2)
    assert not db.try_consume(7, "lessons", 2)
    db.refund(7, "lessons")
    assert db.try_consume(7, "lessons", 2)
    assert db.usage_today(7) == (2, 0)


def test_quota_new_day(db, monkeypatch):
    from edu import db as dbmod
    assert db.try_consume(8, "images", 1)
    assert not db.try_consume(8, "images", 1)
    monkeypatch.setattr(dbmod, "today", lambda: "2099-01-01")
    assert db.try_consume(8, "images", 1)


def test_parallel_quota_is_atomic(db):
    import threading
    ok = []

    def worker():
        ok.append(db.try_consume(9, "lessons", 3))

    threads = [threading.Thread(target=worker) for _ in range(10)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    assert ok.count(True) == 3


def test_catalog_upsert_states(db, bundled_result):
    lesson, raw = bundled_result.lessons[0], bundled_result.raws[0]
    meta = lesson["meta"]
    key = miya.topic_key(meta["topic"])
    assert db.upsert_catalog(meta["grade"], meta["pair"], meta["topic"], key, 45, lesson, raw, "git") == "new"
    assert db.upsert_catalog(meta["grade"], meta["pair"], meta["topic"], key, 45, lesson, raw, "git") == "same"
    row = db.find_catalog(meta["grade"], meta["pair"], key)
    db.update_catalog(row["id"], image_file_id="IMG", card_file_id="CARD", voice_file_ids={"text": "V"})
    raw2 = dict(raw, Uyga_vazifa="Yangi uyga vazifa")
    lesson2 = dict(lesson, homework="Yangi uyga vazifa")
    assert db.upsert_catalog(meta["grade"], meta["pair"], meta["topic"], key, 45, lesson2, raw2, "import") == "updated"
    row = db.get_catalog(row["id"])
    assert row["image_file_id"] == "IMG" and row["card_file_id"] is None and row["voice_file_ids"] == {}


def test_settings_and_users(db):
    db.set_setting("daily_lessons", "9")
    assert db.get_setting("daily_lessons") == "9"
    db.touch_user(1, "Aziz", "aziz")
    db.touch_user(1, "Aziz Ahmad", "aziz")
    assert db.user_name(1) == "Aziz Ahmad (@aziz)"
    assert db.set_blocked(1, True) and db.is_blocked(1)
    assert db.user_ids() == []


# ───────── AI ulanishi ─────────

def test_openai_drops_unsupported_param_and_falls_back_to_json_object(tmp_path, sample_lesson):
    seen = []

    def handler(request):
        if request.url.path == "/v1/models":
            return httpx.Response(200, json={"data": [{"id": "gpt-5.4-mini"}]})
        body = json.loads(request.content)
        seen.append(body)
        if "reasoning" in body:
            return httpx.Response(400, json={"error": {"message": "Unsupported parameter: 'reasoning.effort'",
                                                       "param": "reasoning.effort", "code": "unsupported_parameter"}})
        if body["text"]["format"]["type"] == "json_schema":
            return httpx.Response(400, json={"error": {"message": "Invalid schema for response_format 'lesson'"}})
        return httpx.Response(200, json={"status": "completed", "output": [{"type": "message", "content": [
            {"type": "output_text", "text": json.dumps(sample_lesson)}]}]})

    ai = AIService(make_cfg(tmp_path), transport=httpx.MockTransport(handler))

    async def go():
        try:
            return await ai.lesson("sys", "user", prompts.lesson_schema())
        finally:
            await ai.close()

    data, model = asyncio.run(go())
    assert model == "gpt-5.4-mini" and data["title"] == sample_lesson["title"]
    assert "reasoning" not in seen[-1] and seen[-1]["text"]["format"]["type"] == "json_object"


def test_openai_model_not_found_tries_next(tmp_path, sample_lesson):
    def handler(request):
        if request.url.path == "/v1/models":
            raise httpx.ConnectError("no list")
        body = json.loads(request.content)
        if body["model"] == "gpt-5.4-mini":
            return httpx.Response(404, json={"error": {"message": "The model `gpt-5.4-mini` does not exist",
                                                       "code": "model_not_found"}})
        return httpx.Response(200, json={"status": "completed", "output": [{"type": "message", "content": [
            {"type": "output_text", "text": json.dumps(sample_lesson)}]}]})

    ai = AIService(make_cfg(tmp_path), transport=httpx.MockTransport(handler))

    async def go():
        try:
            return await ai.lesson("sys", "user", prompts.lesson_schema())
        finally:
            await ai.close()

    _, model = asyncio.run(go())
    assert model == "gpt-6-luna"


@pytest.mark.parametrize("status,code,msg,kind", [
    (401, "", "Incorrect API key", "auth"),
    (429, "insufficient_quota", "You exceeded your current quota", "quota"),
    (400, "moderation_blocked", "Your request was rejected by the safety system", "safety"),
    (500, "", "server error", "server"),
])
def test_openai_error_kinds(tmp_path, status, code, msg, kind, monkeypatch):
    async def no_sleep(*a, **k):
        return None
    monkeypatch.setattr(asyncio, "sleep", no_sleep)
    state = {"fail_status": status, "fail_code": code, "fail_message": msg}
    cfg = make_cfg(tmp_path, text_model="gpt-5.4-mini")
    ai = AIService(cfg, transport=openai_transport(state))

    async def go():
        try:
            await ai.lesson("sys", "user", prompts.lesson_schema())
        finally:
            await ai.close()

    with pytest.raises(AIError) as err:
        asyncio.run(go())
    assert err.value.kind == kind and err.value.user_message


def test_model_preferences_pick_newest(tmp_path):
    from edu.ai import GeminiProvider, OpenAIProvider
    cfg = make_cfg(tmp_path)
    o = OpenAIProvider("k", None, cfg)
    assert o.preferences("image", ["gpt-image-1", "gpt-image-2", "gpt-image-2.5-flare"])[0] == "gpt-image-2"
    assert o.preferences("text", ["gpt-4o-mini", "gpt-5.4-mini", "gpt-7-mini"])[0] == "gpt-7-mini"
    g = GeminiProvider("k", None, cfg)
    assert g.preferences("text", ["gemini-2.5-flash", "gemini-3.5-flash", "gemini-3.8-flash", "gemini-3.5-flash-lite"])[0] == "gemini-3.8-flash"
    assert g.preferences("image", None)[0] == "gemini-3.1-flash-image"


def test_image_prompt_uses_exact_counts(sample_lesson):
    lesson = L.normalize_lesson(sample_lesson, grade=2, pair_key="ona_mat", topic="Bog‘dagi hosil", minutes=45)
    p = prompts.build_image_prompt(lesson, 2)
    assert "exactly 6 red apples in the left basket" in p and "exactly 4 red apples" in p
    assert "no text" in p.lower()


# ───────── kartalar va Word ─────────

def test_count_card_draws_exact_numbers():
    groups = [{"label": "1-savat", "object_key": "apple", "count": 6},
              {"label": "2-savat", "object_key": "pear", "count": 13},
              {"label": "Ko‘p", "object_key": "egg", "count": 360}]
    png, drawn = card.render_count_card("Sinov", groups, "footer")
    assert drawn == [6, 13, 1]
    assert Image.open(io.BytesIO(png)).width == 1600


def test_table_card_and_lesson_card(bundled_result):
    for lesson in bundled_result.lessons[:40]:
        png = card.render_lesson_card(lesson, "f")
        assert png and Image.open(io.BytesIO(png)).size[0] == 1600


def test_all_objects_have_images():
    missing = [k for k in objects.keys() if not objects.image_path(k)]
    assert not missing


def test_docx_export(bundled_result, sample_lesson):
    ready = bundled_result.lessons[0]
    data = docx_export.build_docx(ready, image=None, card=card.render_lesson_card(ready, "f"))
    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "DARS ISHLANMASI" in text and "Javoblar (o‘qituvchi uchun)" in text
    assert len(doc.inline_shapes) == 1 and len(doc.tables) == 2
    ai_lesson = L.normalize_lesson(sample_lesson, grade=2, pair_key="ona_mat", topic="Bog‘dagi hosil", minutes=45)
    buf = io.BytesIO()
    Image.new("RGB", (300, 200), (100, 150, 200)).save(buf, "JPEG")
    data = docx_export.build_docx(ai_lesson, image=buf.getvalue(), card=card.render_lesson_card(ai_lesson, "f"))
    assert len(Document(io.BytesIO(data)).inline_shapes) == 2
    assert docx_export.file_name(ai_lesson) == "Dars_2-sinf_Bogdagi_hosil.docx"


# ───────── ovoz ─────────

def test_tts_prepare_rules():
    s = tts.prepare("1-sinf. 6+4=10. Uzunlik 18 cm. 1–3-topshiriqlar. Soat 14:00 da. «Salom» 🍎")
    assert "birinchi sinf" in s and "6 qoʻshuv 4 teng 10" in s and "18 santimetr" in s
    assert "1 dan 3 gacha" in s and "soat 14" in s and "🍎" not in s and "«" not in s


def test_tts_retries_and_switches_voice(monkeypatch):
    calls = []

    async def flaky(text, voice, timeout):
        calls.append(voice)
        if voice == tts.VOICES["madina"]:
            raise RuntimeError("xizmat javob bermadi")
        return b"MP3"

    async def no_sleep(*a, **k):
        return None

    monkeypatch.setattr(tts, "_speak", flaky)
    monkeypatch.setattr(asyncio, "sleep", no_sleep)
    audio = asyncio.run(tts.synthesize("Salom bolalar. Bugun dars.", tts.VOICES["madina"]))
    assert audio == b"MP3" and calls == [tts.VOICES["madina"], tts.VOICES["madina"], tts.VOICES["sardor"]]


def test_tts_long_text_is_chunked(monkeypatch):
    parts = []

    async def ok(text, voice, timeout):
        parts.append(text)
        return b"x"

    monkeypatch.setattr(tts, "_speak", ok)
    asyncio.run(tts.synthesize("Bu uzun gap. " * 250))
    assert len(parts) > 1 and all(len(p) <= tts.CHUNK_CHARS for p in parts)


# ───────── sozlamalar ─────────

def test_load_config(monkeypatch):
    for k in ("OPENAI_API_KEY", "GEMINI_API_KEY", "GOOGLE_API_KEY", "AI_PROVIDER"):
        monkeypatch.delenv(k, raising=False)
    monkeypatch.setenv("TOKEN", " 123:abc ")
    monkeypatch.setenv("ADMIN_ID", "42")
    monkeypatch.setenv("ADMIN_IDS", "43, 44")
    monkeypatch.setenv("GEMINI_API_KEY", "g")
    monkeypatch.setenv("EDU_DAILY_LESSONS", "abc")
    cfg = load_config()
    assert cfg.token == "123:abc" and cfg.admin_ids == {42, 43, 44}
    assert cfg.text_provider == "gemini" and cfg.daily_lessons == 5
    monkeypatch.setenv("OPENAI_API_KEY", "o")
    monkeypatch.setenv("AI_PROVIDER", "gemini")
    assert load_config().text_provider == "gemini"
    monkeypatch.delenv("AI_PROVIDER")
    assert load_config().text_provider == "openai"
