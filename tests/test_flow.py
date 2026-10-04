"""
To‘liq oqim sinovi: soxta Telegram + soxta AI bilan botning hamma asosiy yo‘llari.
Haqiqiy internet, token va API kerak emas.
"""
import asyncio
import io

import openpyxl
import pytest
from telegram import Update

from edu import miya, tts
from edu.ai import AIService
from edu.app import build_application

from .conftest import make_cfg
from .fake_ai import gemini_transport, openai_transport
from .fake_telegram import (FakeTelegram, callback_update, document_update, photo_update, text_update,
                            voice_update)

TEACHER, OTHER, ADMIN = 5001, 5002, 1


async def fake_tts(text, voice="uz-UZ-MadinaNeural", timeout=45):
    fake_tts.calls.append((text, voice))
    return b"ID3-fake-mp3"


fake_tts.calls = []


@pytest.fixture
async def bot(tmp_path, seeded_db, sample_lesson, monkeypatch):
    monkeypatch.setattr(tts, "synthesize", fake_tts)
    fake_tts.calls.clear()
    state = {"lesson": sample_lesson}
    cfg = make_cfg(tmp_path)
    ai = AIService(cfg, transport=openai_transport(state))
    tg = FakeTelegram()
    app = build_application(cfg, db=seeded_db, ai=ai, request=tg, rate_limit=False, seed_miya=False)
    await app.initialize()

    async def send(update_dict):
        await app.process_update(Update.de_json(update_dict, app.bot))

    async def settle():
        for _ in range(50):
            jobs = list(app.bot_data["jobs"])
            if not jobs:
                return
            await asyncio.gather(*jobs, return_exceptions=True)

    yield {"app": app, "tg": tg, "send": send, "settle": settle, "state": state, "db": seeded_db, "cfg": cfg}
    await settle()
    await app.shutdown()
    await ai.close()


def markup_buttons(params):
    rm = params.get("reply_markup")
    if not rm:
        return []
    if isinstance(rm, str):
        import json
        rm = json.loads(rm)
    rows = rm.get("inline_keyboard") or rm.get("keyboard") or []
    out = []
    for row in rows:
        for b in row:
            out.append((b.get("text") if isinstance(b, dict) else b, b.get("callback_data") if isinstance(b, dict) else None))
    return out


def find_cb(params, startswith):
    for text, data in markup_buttons(params):
        if data and data.startswith(startswith):
            return data
    raise AssertionError(f"{startswith} tugmasi yo‘q: {markup_buttons(params)}")


async def test_start_and_help(bot):
    await bot["send"](text_update(TEACHER, "/start"))
    params, _ = bot["tg"].last("sendMessage")
    assert "216 ta tayyor dars" in params["text"]
    assert any(t == "📚 Dars tayyorlash" for t, _ in markup_buttons(params))
    await bot["send"](text_update(TEACHER, "ℹ️ Yordam"))
    params, _ = bot["tg"].last("sendMessage")
    assert "3 qadam" in params["text"]
    sections = [d for _, d in markup_buttons(params) if d]
    assert "h:ready" in sections and "h:admin" not in sections
    for key in ("ready", "ai", "tools", "voice", "pics", "faq"):
        await bot["send"](callback_update(TEACHER, f"h:{key}"))
        p, _ = bot["tg"].last("editMessageText")
        assert len(p["text"]) > 80
    # admin bo‘limini oddiy foydalanuvchi ko‘rmaydi
    await bot["send"](callback_update(TEACHER, "h:admin"))
    p, _ = bot["tg"].last("editMessageText")
    assert "Admin uchun" not in p["text"]
    await bot["send"](text_update(ADMIN, "/yordam"))
    params, _ = bot["tg"].last("sendMessage")
    assert "h:admin" in [d for _, d in markup_buttons(params)]


async def test_ready_lesson_flow(bot):
    tg, send = bot["tg"], bot["send"]
    await send(text_update(TEACHER, "📚 Dars tayyorlash"))
    params, _ = tg.last("sendMessage")
    assert "Sinfni tanlang" in params["text"]
    await send(callback_update(TEACHER, "w:g:1"))
    await send(callback_update(TEACHER, "w:p:ona_mat"))
    params, _ = tg.last("editMessageText")
    ready = [(t, d) for t, d in markup_buttons(params) if d and d.startswith("w:c:")]
    assert len(ready) == 6, ready
    assert any("Bog‘dagi hosil" in t for t, _ in ready)
    cid = [d for t, d in ready if "Bog‘dagi hosil" in t][0]
    ai_calls_before = len(bot["state"].get("calls", []))
    tg.clear()
    await send(callback_update(TEACHER, cid))
    methods = tg.methods()
    assert methods.count("sendPhoto") == 1          # ko‘rgazma kartasi
    assert methods.count("sendMessage") >= 3
    final, _ = tg.last("sendMessage")
    assert "Dars tayyor" in final["text"]
    assert len(bot["state"].get("calls", [])) == ai_calls_before, "tayyor dars AI chaqirmasligi kerak"
    catalog_id = int(cid.split(":")[2])
    row = bot["db"].get_catalog(catalog_id)
    assert row["card_file_id"] and row["views"] == 1

    # javoblar
    tg.clear()
    await send(callback_update(TEACHER, f"x:c:ans:{catalog_id}"))
    text = tg.last("sendMessage")[0]["text"]
    assert "Javoblar" in text and "✅" in text
    # Word
    await send(callback_update(TEACHER, f"x:c:doc:{catalog_id}"))
    params, files = tg.last("sendDocument")
    assert files, "Word fayl yuborilmadi"
    # ovoz: birinchi marta tayyorlanadi, ikkinchisida keshdan
    await send(callback_update(TEACHER, f"x:c:tts:{catalog_id}:text"))
    await send(callback_update(TEACHER, f"x:c:tts:{catalog_id}:text"))
    await send(callback_update(TEACHER, f"x:c:tts:{catalog_id}:tasks"))
    assert len(fake_tts.calls) == 2, fake_tts.calls
    assert "Birinchi topshiriq" in fake_tts.calls[1][0]
    assert bot["tg"].methods().count("sendVoice") == 3

    # AI rasm: birinchi o‘qituvchi chizdiradi, ikkinchisi keshdan bepul oladi
    await send(callback_update(TEACHER, f"x:c:img:{catalog_id}"))
    await bot["settle"]()
    row = bot["db"].get_catalog(catalog_id)
    assert row["image_file_id"] and row["image_note"].startswith("✅")
    images_before = sum(1 for c in bot["state"]["calls"] if c[1] == "/v1/images/generations")
    await send(callback_update(OTHER, f"x:c:img:{catalog_id}"))
    await bot["settle"]()
    images_after = sum(1 for c in bot["state"]["calls"] if c[1] == "/v1/images/generations")
    assert images_after == images_before
    # baho
    await send(callback_update(TEACHER, f"x:c:rate:{catalog_id}:1"))
    assert bot["db"].stats()["likes"] == 1


async def test_free_text_finds_ready_lessons(bot):
    await bot["send"](text_update(TEACHER, "bog'dagi hosil"))
    params, _ = bot["tg"].last("sendMessage")
    ready = [d for _, d in markup_buttons(params) if d and d.startswith("w:c:")]
    assert len(ready) == 4  # 1–4-sinflar


async def test_ai_lesson_flow_and_limits(bot):
    tg, send = bot["tg"], bot["send"]
    await send(text_update(TEACHER, "Qishki o'yinlar"))
    params, _ = tg.last("sendMessage")
    assert find_cb(params, "w:topic") == "w:topic"
    await send(callback_update(TEACHER, "w:topic"))
    await send(callback_update(TEACHER, "w:g:2"))
    await send(callback_update(TEACHER, "w:p:ona_mat"))
    params, _ = tg.last("editMessageText")
    assert "AI bilan yangi dars" in params["text"]
    await send(callback_update(TEACHER, "w:m:90"))
    await send(callback_update(TEACHER, "w:extra"))
    await send(text_update(TEACHER, "guruhda ishlash bo'lsin"))
    params, _ = tg.last("sendMessage")
    assert "guruhda ishlash" in params["text"] and "90 daqiqa" in params["text"]
    tg.clear()
    await send(callback_update(TEACHER, "w:go"))
    await bot["settle"]()
    methods = tg.methods()
    assert methods.count("sendPhoto") == 2  # hisob kartasi + AI rasm
    final, _ = tg.last("sendMessage")
    assert "Dars tayyor" in final["text"] and "5-topshiriq" in final["text"]  # hisob farqi ogohlantirishi
    rows = bot["db"].list_lessons(TEACHER)
    assert len(rows) == 1 and rows[0]["topic"] == "Qishki o‘yinlar"
    lesson_row = bot["db"].get_lesson(rows[0]["id"])
    assert lesson_row["minutes"] == 90 and sum(s["minutes"] for s in lesson_row["lesson"]["lesson_plan"]) == 90
    assert lesson_row["image_file_id"] and lesson_row["card_file_id"]
    req = [b for b in bot["state"]["bodies"] if (b.get("text") or {}).get("format", {}).get("name") == "lesson"][-1]
    assert "Qishki o‘yinlar" in req["input"] and "guruhda ishlash" in req["input"]
    assert req["reasoning"] == {"effort": "low"}
    assert bot["db"].usage_today(TEACHER) == (1, 1)

    # arxiv → ochish → o‘chirish
    await send(text_update(TEACHER, "🗂 Mening darslarim"))
    params, _ = tg.last("sendMessage")
    open_cb = find_cb(params, "x:l:open:")
    await send(callback_update(TEACHER, open_cb))
    await send(callback_update(OTHER, open_cb))  # begona foydalanuvchi ochololmaydi
    assert tg.last("answerCallbackQuery")[0].get("show_alert") is True
    lid = int(open_cb.split(":")[3])
    await send(callback_update(TEACHER, f"x:l:delok:{lid}"))
    assert bot["db"].count_lessons(TEACHER) == 0

    # limit: kuniga 2 ta dars
    for _ in range(2):
        await send(text_update(TEACHER, "Yangi mavzu sinov"))
        await send(callback_update(TEACHER, "w:topic"))
        await send(callback_update(TEACHER, "w:g:3"))
        await send(callback_update(TEACHER, "w:p:tab_mat"))
        await send(callback_update(TEACHER, "w:img:0"))
        await send(callback_update(TEACHER, "w:go"))
        await bot["settle"]()
    params, _ = tg.last("sendMessage")
    assert "limiti tugadi" in params["text"]
    assert bot["db"].usage_today(TEACHER)[0] == 2


async def test_ai_failure_refunds_and_notifies(bot):
    bot["state"]["fail_status"] = 429
    bot["state"]["fail_code"] = "insufficient_quota"
    bot["state"]["fail_message"] = "You exceeded your current quota"
    send, tg = bot["send"], bot["tg"]
    await send(text_update(TEACHER, "Muz va suv"))
    await send(callback_update(TEACHER, "w:topic"))
    await send(callback_update(TEACHER, "w:g:1"))
    await send(callback_update(TEACHER, "w:p:ona_tab"))
    await send(callback_update(TEACHER, "w:go"))
    await bot["settle"]()
    texts = tg.texts()
    assert any("Dars tayyorlanmadi" in t and "mablag‘" in t for t in texts)
    assert bot["db"].usage_today(TEACHER) == (0, 0)  # limit qaytarildi
    admin_msgs = [p for a, p, _ in tg.calls if a == "sendMessage" and str(p.get("chat_id")) == str(ADMIN)]
    assert any("AI xatosi" in p["text"] for p in admin_msgs)


async def test_voice_message_topic(bot):
    bot["state"]["transcript"] = "Боғдаги ҳосил"  # kirill — lotinga o‘giriladi
    await bot["send"](voice_update(TEACHER))
    texts = bot["tg"].texts()
    assert any("Eshitganim: «Bog‘dagi hosil»" in t for t in texts)
    params, _ = bot["tg"].last("sendMessage")
    assert len([d for _, d in markup_buttons(params) if d and d.startswith("w:c:")]) == 4
    # juda uzun ovoz
    await bot["send"](voice_update(TEACHER, duration=300))
    assert "1,5 daqiqadan" in bot["tg"].last("sendMessage")[0]["text"]


async def test_admin_panel_broadcast_settings_and_miya(bot, tmp_path):
    tg, send, db = bot["tg"], bot["send"], bot["db"]
    await send(text_update(TEACHER, "/start"))
    await send(text_update(OTHER, "/start"))
    await send(text_update(ADMIN, "/admin"))
    params, _ = tg.last("sendMessage")
    assert "Admin panel" in params["text"]
    await send(text_update(TEACHER, "/admin"))  # oddiy foydalanuvchiga panel chiqmaydi
    await send(callback_update(ADMIN, "adm:stats"))
    assert "STATISTIKA" in tg.last("editMessageText")[0]["text"]
    await send(callback_update(TEACHER, "adm:stats"))
    assert tg.last("answerCallbackQuery")[0].get("show_alert") is True

    # sozlamalar
    await send(callback_update(ADMIN, "adm:set:lessons:1"))
    await send(callback_update(ADMIN, "adm:set:img"))
    await send(callback_update(ADMIN, "adm:set:voice"))
    assert db.get_setting("daily_lessons") == "3" and db.get_setting("images_enabled") == "0"
    assert db.get_setting("tts_voice") == "sardor"
    await send(callback_update(ADMIN, "adm:set:img"))

    # AI holati
    await send(callback_update(ADMIN, "adm:ai"))
    assert "gpt-5.4-mini" in tg.last("editMessageText")[0]["text"]

    # reklama: rasm yuboriladi → tasdiq → hammaga nusxa
    await send(callback_update(ADMIN, "adm:bc"))
    await send(photo_update(ADMIN))
    assert "yuborilsinmi" in tg.last("sendMessage")[0]["text"]
    await send(callback_update(ADMIN, "adm:bc:go"))
    await bot["settle"]()
    copies = [p for a, p, _ in tg.calls if a == "copyMessage"]
    assert len(copies) == 3
    assert "Reklama tugadi" in tg.last("editMessageText")[0]["text"]

    # miya eksport → o‘sha faylni import (o‘zgarishsiz qaytadi)
    await send(text_update(ADMIN, "/miya_export"))
    _, files = tg.last("sendDocument")
    data = list(files.values())[0][1]
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb["Miya"].max_row == 217
    result = miya.parse_workbook(data)
    assert len(result.lessons) == 216 and not result.errors

    # admin AI darsini miyaga qo‘shadi
    lesson_id = db.save_lesson(TEACHER, 3, "oqish_tab", "Bulutlar", "", 45,
                               {"title": "Bulutlar", "meta": {"grade": 3, "pair": "oqish_tab", "topic": "Bulutlar",
                                                              "minutes": 45}, "tasks": [], "origin": "ai",
                                "image": {"count_groups": []}}, "test")
    await send(callback_update(ADMIN, f"adm:tomiya:{lesson_id}"))
    assert db.count_catalog() == 217


async def test_miya_import_via_document(bot, monkeypatch):
    rows_before = bot["db"].count_catalog()
    with open(miya.BUNDLED_PATH, "rb") as fh:
        content = fh.read()
    wb = openpyxl.load_workbook(io.BytesIO(content))
    ws = wb["Miya"]
    header = [c.value for c in ws[1]]
    new_row = [c.value for c in ws[2]]
    new_row[header.index("Mavzu")] = "Yangi sinov mavzusi"
    ws.append(new_row)
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()

    from telegram import File

    async def fake_download(self, *a, **k):
        return bytearray(data)

    monkeypatch.setattr(File, "download_as_bytearray", fake_download)
    await bot["send"](text_update(ADMIN, "/miya_import"))
    await bot["send"](document_update(ADMIN, "Fanlar_Miya.xlsx", len(data)))
    text = bot["tg"].last("editMessageText")[0]["text"]
    assert "yangi: 1" in text
    await bot["send"](callback_update(ADMIN, "adm:miya:ok"))
    assert bot["db"].count_catalog() == rows_before + 1
    # oddiy foydalanuvchi fayli qabul qilinmaydi
    await bot["send"](document_update(TEACHER, "x.xlsx"))
    assert "faylni qabul qilmayman" in bot["tg"].last("sendMessage")[0]["text"]


async def test_blocked_user_is_ignored(bot):
    await bot["send"](text_update(OTHER, "/start"))
    await bot["send"](text_update(ADMIN, f"/block {OTHER}"))
    bot["tg"].clear()
    await bot["send"](text_update(OTHER, "📚 Dars tayyorlash"))
    assert bot["tg"].methods() == []
    await bot["send"](text_update(ADMIN, f"/unblock {OTHER}"))
    await bot["send"](text_update(OTHER, "📚 Dars tayyorlash"))
    assert "sendMessage" in bot["tg"].methods()


async def test_stale_and_unknown_buttons(bot):
    await bot["send"](callback_update(TEACHER, "w:g:2"))   # bosqichsiz eski tugma → yangi bosqich
    assert "Sinfni tanlang" in bot["tg"].last("editMessageText")[0]["text"]
    await bot["send"](callback_update(TEACHER, "add_15"))  # kiyim botidan qolgan tugma
    assert "eskirgan" in bot["tg"].last("answerCallbackQuery")[0]["text"]
    await bot["send"](callback_update(TEACHER, "x:c:ans:999999"))
    assert "topilmadi" in bot["tg"].last("answerCallbackQuery")[0]["text"]


async def test_works_without_ai_key(tmp_path, seeded_db, monkeypatch):
    monkeypatch.setattr(tts, "synthesize", fake_tts)
    cfg = make_cfg(tmp_path, openai_key="", text_provider="", image_provider="")
    ai = AIService(cfg)
    tg = FakeTelegram()
    app = build_application(cfg, db=seeded_db, ai=ai, request=tg, rate_limit=False, seed_miya=False)
    await app.initialize()
    send = lambda u: app.process_update(Update.de_json(u, app.bot))
    await send(text_update(TEACHER, "/start"))
    await send(text_update(TEACHER, "📚 Dars tayyorlash"))
    await send(callback_update(TEACHER, "w:g:4"))
    await send(callback_update(TEACHER, "w:p:oqish_tex"))
    params, _ = tg.last("editMessageText")
    ready = [d for _, d in markup_buttons(params) if d and d.startswith("w:c:")]
    assert len(ready) == 6
    await send(callback_update(TEACHER, ready[0]))
    final, _ = tg.last("sendMessage")
    assert "Dars tayyor" in final["text"]
    assert not [d for _, d in markup_buttons(final) if d and ":img:" in d]  # AI rasm tugmasi yo‘q
    await send(text_update(TEACHER, "Mutlaqo yangi mavzu"))
    assert "tayyor dars topilmadi" in tg.last("sendMessage")[0]["text"]
    await send(voice_update(TEACHER))
    assert "AI ulanmagan" in tg.last("sendMessage")[0]["text"]
    await app.shutdown()
    await ai.close()


async def test_gemini_provider_flow(tmp_path, seeded_db, sample_lesson, monkeypatch):
    monkeypatch.setattr(tts, "synthesize", fake_tts)
    state = {"lesson": sample_lesson}
    cfg = make_cfg(tmp_path, openai_key="", gemini_key="AIza-test", text_provider="gemini", image_provider="gemini")
    ai = AIService(cfg, transport=gemini_transport(state))
    tg = FakeTelegram()
    app = build_application(cfg, db=seeded_db, ai=ai, request=tg, rate_limit=False, seed_miya=False)
    await app.initialize()
    send = lambda u: app.process_update(Update.de_json(u, app.bot))
    await send(text_update(TEACHER, "Qishki tabiat sirlari"))
    await send(callback_update(TEACHER, "w:topic"))
    await send(callback_update(TEACHER, "w:g:2"))
    await send(callback_update(TEACHER, "w:p:ona_mat"))
    await send(callback_update(TEACHER, "w:go"))
    for _ in range(20):
        jobs = list(app.bot_data["jobs"])
        if not jobs:
            break
        await asyncio.gather(*jobs, return_exceptions=True)
    assert "Dars tayyor" in tg.last("sendMessage")[0]["text"]
    paths = [p for p, _ in state["bodies"]]
    assert any("gemini-3.5-flash:generateContent" in p for p in paths)
    assert any("gemini-3.1-flash-image:generateContent" in p for p in paths)
    lesson_body = [b for p, b in state["bodies"] if "systemInstruction" in b][0]
    assert lesson_body["generationConfig"]["responseMimeType"] == "application/json"
    await app.shutdown()
    await ai.close()
